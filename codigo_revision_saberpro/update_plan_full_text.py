# -*- coding: utf-8 -*-
"""Reemplaza la columna de comentario resumido por el texto completo del
comentario original (traducido al español, sin resumir)."""
import openpyxl
from openpyxl.styles import Alignment

wb = openpyxl.load_workbook("/home/claude/plan_revision_saberpro.xlsx")
ws = wb["Plan de revisión"]

header_row = 4

# Cambiar encabezado de la columna D
ws.cell(row=header_row, column=4, value="Comentario completo (texto íntegro, traducido)")

full_comments = {
    "E1": "Existe la necesidad de una justificación y reproducibilidad más sólidas de la solución de clustering, incluyendo la selección de ocho clústeres. En este sentido, debe proporcionarse una documentación más clara de la configuración analítica.",
    "E2": "Además, debe proporcionarse evidencia más sólida de estabilidad respecto al procedimiento completo UMAP–K-means, en contraposición a evaluar únicamente los efectos de la inicialización de K-means mientras se mantiene fija la representación de UMAP.",
    "E3": "En segundo lugar, debe clarificarse la relación entre los perfiles descriptivos, su validación y las aplicaciones propuestas. En particular, será importante profundizar en cómo las características particulares de los perfiles podrían informar decisiones institucionales orientadas a la equidad.",

    "R1-1": "Sería deseable ampliar (o complementar) los detalles operativos, particularmente en lo referente a los hiperparámetros de UMAP, la inicialización del clustering y los criterios de consistencia, y cómo se controla la variabilidad del proceso (p. ej., semillas aleatorias, sensibilidad del clustering y estabilidad entre corridas).",
    "R1-2": "Aunque se menciona la selección de K=8, el soporte empírico para esta elección podría presentarse con mayor claridad, incluyendo los criterios utilizados (p. ej., el método del codo, el índice de Davies–Bouldin, u otras medidas de validez de clústeres).",
    "R1-3": "Dado que el objetivo es la “toma de decisiones orientada a la equidad”, sería valioso explicar de manera más explícita cómo los perfiles identificados se traducen en planes de intervención y qué variables o patrones desencadenan decisiones específicas (p. ej., estrategias de priorización o recomendaciones adaptadas a territorios o grupos poblacionales particulares).",

    "R2-1": "La estructura de clústeres puede depender fuertemente de UMAP. Los valores de Silhouette de 0.05 para el espacio original y 0.09 para K-Means sin reducción indican una separación débil bajo la codificación elegida y la métrica euclidiana, aunque no descartan por completo la existencia de estructura no lineal. Aun cuando el aumento a 0.43 demuestra separación en el embedding de UMAP, esto no establece que los datos originales contengan ocho clústeres naturales ni confirma una variedad (manifold) no lineal. Para abordar este problema, los autores pueden calibrar el pipeline completo y ajustado frente a conjuntos de datos nulos repetidos. Si la solución observada no supera claramente los puntos de referencia nulos, la separación aparente podría reflejar el procedimiento de reducción de dimensionalidad y selección de modelo, en lugar de una estructura de clústeres significativa en los datos originales.",
    "R2-2": "Evidencia de estabilidad. Un ARI medio de 0.69 (DE = 0.14) indica un acuerdo significativo pero imperfecto. Sería muy valioso que el estudio reportara explícitamente cómo se calculó el ARI, su distribución completa, la estabilidad específica por clúster, y una matriz de consenso o co-clustering. Además, el análisis de robustez actual varía únicamente la inicialización de K-Means, manteniendo fijo el embedding de UMAP. Para evaluar la estabilidad del pipeline completo, los autores deberían repetir ambas etapas utilizando diferentes semillas aleatorias de UMAP y muestras de ajuste de 80,000 observaciones extraídas de forma independiente, y luego evaluar el acuerdo entre las asignaciones de clúster resultantes.",
    "R2-3": "Comparación entre algoritmos de clustering. Los algoritmos se evaluaron sobre espacios de búsqueda diferentes, dado que DBSCAN se restringió a dos dimensiones, mientras que los valores de min_cluster_size de HDBSCAN son muy pequeños en relación con 200,000 observaciones. No se reportan decisiones relevantes de GMM, incluyendo la especificación de la covarianza y la inicialización. Los autores podrían evaluar dimensionalidades comparables y rangos de hiperparámetros más amplios antes de concluir que los métodos basados en densidad no son adecuados o que los datos carecen de una estructura de clústeres de densidad variable.",
    "R2-4": "Datos faltantes. La educación materna y paterna presentan aproximadamente 22-23% de valores faltantes, y ambas se reemplazan por un único valor de mediana. Esto puede crear concentraciones artificiales, atenuar la variación y alterar los vecindarios de UMAP. La verificación de casos para los Clústeres 1 y 5 es útil, pero no evalúa los efectos sobre la partición completa. Por favor repitan el análisis utilizando imputación múltiple, indicadores de valor faltante, u otra estrategia principiada, y comparen las asignaciones de clúster.",
    "R2-5": "Diferencias de medias reportadas y construcción del clustering. El logro académico y las variables socioeconómicas se utilizan para crear los clústeres, por lo que las diferencias en sus medias por clúster y las pruebas asociadas no son hallazgos independientes. Con n = 452,020, incluso diferencias pequeñas serán estadísticamente significativas. Estos resultados deberían presentarse de manera descriptiva con tamaños de efecto, no como validación. La validación independiente requiere variables que hayan sido excluidas de la construcción del clustering.",
    "R2-6": "Contribución del clustering más allá de análisis más simples. El hallazgo principal de que los Clústeres 0 y 6 tienen medias similares pero restricciones diferentes equivale a una tabulación cruzada de puntajes según horas trabajadas y acceso a internet. Para un público general, el artículo debería mostrar que la pertenencia a un perfil aporta información más allá de la combinación lineal de los insumos, por ejemplo, prediciendo un resultado (puntaje, o idealmente graduación) condicionado en las covariables.",
    "R2-7": "Análisis geográfico y validación externa. Cada fila de la Figura 3 suma aproximadamente 100%, por lo tanto las celdas reportan el porcentaje de examinados dentro de cada departamento asignados a cada clúster, P(clúster | departamento). Varias afirmaciones parecen interpretarlas, en cambio, como P(departamento | clúster). En particular, el 16% reportado para Bogotá y el 14% para Valle del Cauca corresponden al Clúster 4 en la figura, mientras que los valores mostrados para el Clúster 7 son 2% y 1%, respectivamente. Por ejemplo, la afirmación de que el Clúster 7 se concentra en estos centros urbanos no está, por lo tanto, respaldada por la figura tal como se presenta.",
    "R2-8": "Análisis temporal. El análisis combina las convocatorias desde 2021-2 hasta 2023-1, pero no muestra que las escalas de puntaje, las distribuciones de variables y la prevalencia de los perfiles sean comparables entre periodos. Por favor documenten la comparabilidad y reporten la composición de los clústeres por convocatoria. Ajustar el pipeline por separado para cada periodo proporcionaría una prueba de replicación temporal importante y revelaría si los perfiles son estables en lugar de específicos de una cohorte.",
    "R2-9": "Muestra. La similitud entre una submuestra de 200,000 registros y el conjunto de datos del cual fue extraída no establece representatividad nacional. Los autores pueden describir la cobertura, las exclusiones y la posible selección en la participación de Saber Pro. Asimismo, aclarar si las 452,020 observaciones representan estudiantes únicos o registros de examen que puedan incluir presentadores repetidos.",
    "R2-10": "Concentración a nivel institucional y de programa como explicaciones alternativas plausibles. Los estudiantes están anidados dentro de programas e instituciones, pero estas estructuras no se examinan en el estudio ni se mencionan como una limitación. La agregación geográfica no demuestra que los perfiles sean independientes de un número reducido de instituciones o campos de estudio. Como mínimo, los autores pueden reportar la composición institucional y de programa por clúster, y realizar un análisis de sensibilidad de tipo leave-institution-out, leave-program-out, o comparable. Este aspecto es importante para interpretar los perfiles como configuraciones a nivel de estudiante y no institucionales.",
    "R2-11": "El Clúster 7 requiere un examen más detallado. El Clúster 7 contiene solo el 1.5% de las observaciones y forma dos islas desconectadas en la representación de UMAP. Sería muy valioso que los autores evaluaran si esto refleja subpoblaciones reproducibles, valores extremos, o errores de puntuación o preprocesamiento. Las dos islas podrían caracterizarse por separado.",
    "R2-12": "Comparaciones y sobreinterpretaciones. La comparación entre la brecha de Saber Pro y las brechas de PISA no es válida porque las pruebas utilizan escalas y desviaciones estándar diferentes. Sería preferible reportar diferencias estandarizadas en su lugar. Asimismo, las afirmaciones de la sección 6 sobre restricciones estructurales, mecanismos del mercado laboral, exclusión digital, y canales culturales maternos o paternos deberían identificarse explícitamente como hipótesis y no como hallazgos del análisis transversal.",
    "R2-13": "En las conclusiones, las tres respuestas de intervención propuestas son especulativas, y podrían presentarse como un marco orientado a políticas, pero su efectividad e idoneidad requieren evaluación causal.",
}

max_lines_needed = {}
for row in range(header_row + 1, header_row + 1 + len(full_comments)):
    id_val = ws.cell(row=row, column=1).value
    if id_val in full_comments:
        cell = ws.cell(row=row, column=4, value=full_comments[id_val])
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        # estimar altura de fila según longitud del texto (col D ancho ~46)
        n_chars = len(full_comments[id_val])
        est_lines = max(3, (n_chars // 60) + 1)
        ws.row_dimensions[row].height = max(ws.row_dimensions[row].height or 0, est_lines * 14)

# Ampliar un poco la columna D ya que ahora lleva texto completo
ws.column_dimensions["D"].width = 55

wb.save("/home/claude/plan_revision_saberpro.xlsx")
print("updated")
