# -*- coding: utf-8 -*-
"""Construye el plan de revisión (respuesta a revisores) en Excel."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

FONT_NAME = "Arial"

wb = openpyxl.Workbook()

# ---------------------------------------------------------------
# Datos: fuente, ubicación, comentario, acción, tipo, prioridad, complejidad
# ---------------------------------------------------------------
rows = [
    # Editor
    ("E1", "Editor", "General / Métodos",
     "Necesidad de justificación y reproducibilidad más sólida de la solución de clustering, incluida la selección de ocho clústeres; documentar con más claridad la configuración analítica.",
     "Documentar exhaustivamente hiperparámetros de UMAP, inicialización de K-Means y criterios de selección de K=8 en una tabla dedicada.",
     "Documentación", "Alta", "Baja"),
    ("E2", "Editor", "General / Estabilidad",
     "Se requiere evidencia más sólida de estabilidad de todo el procedimiento UMAP+K-Means, no solo de la inicialización de K-Means manteniendo fija la proyección UMAP.",
     "Ejecutar análisis de estabilidad del pipeline completo variando semillas de UMAP y usando submuestras independientes.",
     "Reanálisis", "Alta", "Alta"),
    ("E3", "Editor", "Discusión / Conclusiones",
     "Clarificar la relación entre los perfiles descriptivos, su validación, y sus aplicaciones propuestas; explicar cómo las características de los perfiles informan decisiones institucionales orientadas a la equidad.",
     "Reescribir la sección de discusión/conclusiones vinculando explícitamente cada perfil con decisiones o intervenciones concretas.",
     "Redacción", "Alta", "Media"),
    # Reviewer 1
    ("R1-1", "Revisor 1", "Métodos / UMAP",
     "Ampliar detalles operativos: hiperparámetros de UMAP, inicialización y criterios de consistencia del clustering, control de variabilidad del proceso (semillas aleatorias, sensibilidad y estabilidad entre corridas).",
     "Añadir tabla completa de hiperparámetros y reportar la estabilidad entre corridas (se solapa con E1/E2).",
     "Documentación", "Alta", "Media"),
    ("R1-2", "Revisor 1", "Métodos / Selección de K",
     "El soporte empírico para K=8 debe presentarse con más claridad (método del codo, índice Davies-Bouldin u otras métricas de validez de clústeres).",
     "Reportar explícitamente los criterios cuantitativos usados para seleccionar K=8 (se solapa con E1).",
     "Documentación", "Alta", "Baja"),
    ("R1-3", "Revisor 1", "Discusión / Aplicación",
     "Explicar cómo los perfiles identificados se traducen en planes de intervención concretos y qué variables o patrones disparan decisiones específicas.",
     "Detallar reglas de decisión/priorización por perfil (se solapa con E3).",
     "Redacción", "Media", "Baja"),
    # Reviewer 2 (13 points)
    ("R2-1", "Revisor 2", "Metodología / UMAP (pt.1)",
     "La estructura de clústeres puede depender fuertemente de UMAP: Silhouette de 0.05 (espacio original) y 0.09 (K-Means sin reducción) indican separación débil; el aumento a 0.43 en UMAP no confirma que existan ocho clústeres naturales en los datos originales.",
     "Calibrar el pipeline completo (UMAP+K-Means) contra datasets nulos/permutados repetidos y comparar la separación observada frente al benchmark nulo.",
     "Reanálisis", "Alta", "Alta"),
    ("R2-2", "Revisor 2", "Estabilidad (pt.2)",
     "ARI medio de 0.69 (DE=0.14) es informativo pero incompleto: falta reportar cómo se calculó, su distribución completa, estabilidad por clúster y una matriz de consenso/co-clustering. La robustez actual solo varía la inicialización de K-Means manteniendo fija la proyección UMAP.",
     "Repetir ambas etapas (UMAP y K-Means) con distintas semillas de UMAP y submuestras independientes de 80,000 observaciones; reportar matriz de consenso.",
     "Reanálisis", "Alta", "Alta"),
    ("R2-3", "Revisor 2", "Comparación de algoritmos (pt.3)",
     "Los algoritmos se evaluaron en espacios de búsqueda distintos (DBSCAN restringido a 2D; min_cluster_size de HDBSCAN muy pequeño respecto a 200,000 obs.); no se reportan opciones relevantes de GMM (covarianza, inicialización).",
     "Igualar dimensionalidad y ampliar el rango de hiperparámetros de DBSCAN/HDBSCAN/GMM antes de concluir que los métodos basados en densidad no son adecuados.",
     "Reanálisis", "Media", "Media"),
    ("R2-4", "Revisor 2", "Datos faltantes (pt.4)",
     "Educación materna/paterna tiene ~22-23% de valores faltantes, imputados con un único valor de mediana; esto puede crear concentraciones artificiales y alterar los vecindarios de UMAP.",
     "Repetir el análisis con imputación múltiple o indicadores de valor faltante y comparar las asignaciones de clúster resultantes.",
     "Reanálisis", "Alta", "Media"),
    ("R2-5", "Revisor 2", "Validación con variables del clustering (pt.5)",
     "Logro académico y variables socioeconómicas se usan para crear los clústeres, por lo que las diferencias de medias entre clústeres no son hallazgos independientes; con n=452,020 incluso diferencias pequeñas serán estadísticamente significativas.",
     "Presentar estas diferencias como descriptivas con tamaños de efecto, no como validación; aclarar explícitamente en el texto.",
     "Redacción / Reencuadre", "Alta", "Baja"),
    ("R2-6", "Revisor 2", "Contribución del clustering (pt.6)",
     "El hallazgo principal (Clústeres 0 y 6 con medias similares pero restricciones distintas) equivale a una tabulación cruzada de puntaje por horas trabajadas y acceso a internet; falta mostrar que la pertenencia a un perfil aporta información más allá de la combinación lineal de los insumos.",
     "Predecir un resultado (puntaje o, idealmente, graduación) condicionado en las covariables, comparando el modelo con y sin la pertenencia al clúster.",
     "Reanálisis", "Media", "Media"),
    ("R2-7", "Revisor 2", "Análisis geográfico – Figura 3 (pt.7)",
     "Cada fila de la Figura 3 suma ~100%, por lo que las celdas son P(clúster | departamento), no P(departamento | clúster). El 16% de Bogotá y 14% de Valle del Cauca corresponden al Clúster 4, no al Clúster 7 (cuyos valores allí son 2% y 1%). La afirmación de que el Clúster 7 se concentra en estos centros urbanos no está respaldada por la figura.",
     "Corregir la interpretación estadística en el texto y revisar todas las afirmaciones basadas en la Figura 3.",
     "Corrección de error", "Alta", "Baja"),
    ("R2-8", "Revisor 2", "Análisis temporal (pt.8)",
     "Se combinan administraciones de 2021-2 a 2023-1 sin mostrar que las escalas de puntaje, distribuciones de variables y prevalencia de perfiles sean comparables entre periodos.",
     "Documentar la comparabilidad entre periodos, reportar la composición de clústeres por administración, y ajustar el pipeline por separado por periodo como prueba de replicación temporal.",
     "Reanálisis", "Media", "Media"),
    ("R2-9", "Revisor 2", "Muestra / representatividad (pt.9)",
     "La similitud entre la submuestra de 200,000 registros y el conjunto completo no establece representatividad nacional; falta aclarar si las 452,020 observaciones son estudiantes únicos o registros de examen (posibles repetidores).",
     "Describir cobertura, exclusiones y posible selección en la participación de Saber Pro; aclarar la unicidad de los estudiantes.",
     "Documentación", "Media", "Baja"),
    ("R2-10", "Revisor 2", "Concentración institucional/programa (pt.10)",
     "Los estudiantes están anidados en programas e instituciones, no examinados en el estudio; la agregación geográfica no demuestra que los perfiles sean independientes de un número reducido de instituciones o programas.",
     "Reportar la composición institucional/de programa por clúster y realizar un análisis de sensibilidad leave-institution-out / leave-program-out.",
     "Reanálisis", "Alta", "Alta"),
    ("R2-11", "Revisor 2", "Clúster 7 (pt.11)",
     "El Clúster 7 contiene solo 1.5% de las observaciones y forma dos islas desconectadas en la representación UMAP.",
     "Evaluar si refleja subpoblaciones reproducibles, valores extremos o errores de preprocesamiento; caracterizar las dos islas por separado.",
     "Reanálisis", "Media", "Media"),
    ("R2-12", "Revisor 2", "Comparaciones y sobreinterpretación (pt.12)",
     "La comparación de la brecha de Saber Pro con las brechas de PISA no es válida porque usan escalas y desviaciones estándar distintas; las afirmaciones de la sección 6 sobre mecanismos estructurales/laborales/exclusión digital deben identificarse como hipótesis, no como hallazgos.",
     "Reportar diferencias estandarizadas en vez de la comparación directa; reencuadrar explícitamente las afirmaciones causales como hipótesis.",
     "Redacción / Reencuadre", "Alta", "Baja"),
    ("R2-13", "Revisor 2", "Conclusiones / intervención (pt.13)",
     "Las tres respuestas de intervención propuestas son especulativas; pueden presentarse como marco orientado a políticas, pero su efectividad requiere evaluación causal.",
     "Reencuadrar explícitamente como marco especulativo/hipotético, no como recomendaciones validadas.",
     "Redacción / Reencuadre", "Media", "Baja"),
]

headers = ["ID", "Fuente", "Ubicación en el manuscrito", "Comentario (resumen)",
           "Acción requerida", "Tipo de acción", "Prioridad", "Complejidad",
           "Estado", "Notas / respuesta a revisores"]

ws = wb.active
ws.title = "Plan de revisión"

# --- estilos ---
header_fill = PatternFill("solid", fgColor="1F4E78")
header_font = Font(name=FONT_NAME, bold=True, color="FFFFFF", size=11)
title_font = Font(name=FONT_NAME, bold=True, size=14, color="1F4E78")
subtitle_font = Font(name=FONT_NAME, italic=True, size=10, color="595959")
base_font = Font(name=FONT_NAME, size=10)
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")
wrap_center = Alignment(wrap_text=True, vertical="top", horizontal="center")

fuente_fill = {
    "Editor": "FCE4D6",
    "Revisor 1": "DDEBF7",
    "Revisor 2": "E2EFDA",
}
prioridad_fill = {
    "Alta": "F8696B",
    "Media": "FFEB84",
    "Baja": "63BE7B",
}

# --- título ---
ws.merge_cells("A1:J1")
ws["A1"] = "Plan de revisión — Structural inequality in higher education performance (UMAP + K-Means, Saber Pro)"
ws["A1"].font = title_font
ws.merge_cells("A2:J2")
ws["A2"] = "Large-scale Assessments in Education — carta del editor, Revisor 1 (3 comentarios) y Revisor 2 (13 comentarios)"
ws["A2"].font = subtitle_font
ws.row_dimensions[1].height = 22
ws.row_dimensions[2].height = 16

header_row = 4
for j, h in enumerate(headers, start=1):
    c = ws.cell(row=header_row, column=j, value=h)
    c.font = header_font
    c.fill = header_fill
    c.alignment = wrap_center
    c.border = border

for i, r in enumerate(rows, start=header_row + 1):
    id_, fuente, ubicacion, comentario, accion, tipo, prioridad, complejidad = r
    values = [id_, fuente, ubicacion, comentario, accion, tipo, prioridad, complejidad, "Pendiente", ""]
    for j, v in enumerate(values, start=1):
        c = ws.cell(row=i, column=j, value=v)
        c.font = base_font
        c.border = border
        c.alignment = wrap
        if j in (1, 6, 7, 8, 9):
            c.alignment = wrap_center
    ws.cell(row=i, column=2).fill = PatternFill("solid", fgColor=fuente_fill[fuente])
    prio_cell = ws.cell(row=i, column=7)
    prio_cell.fill = PatternFill("solid", fgColor=prioridad_fill[prioridad])
    if prioridad == "Alta":
        prio_cell.font = Font(name=FONT_NAME, size=10, color="FFFFFF", bold=True)

last_row = header_row + len(rows)

# --- anchos de columna ---
widths = {"A": 7, "B": 11, "C": 24, "D": 46, "E": 42, "F": 16, "G": 10, "H": 11, "I": 13, "J": 32}
for col, w in widths.items():
    ws.column_dimensions[col].width = w

for i in range(header_row + 1, last_row + 1):
    ws.row_dimensions[i].height = 60

ws.freeze_panes = "A5"
ws.auto_filter.ref = f"A{header_row}:J{last_row}"

# --- validación de datos (listas desplegables) ---
dv_estado = DataValidation(type="list", formula1='"Pendiente,En progreso,Completado,No se abordará"', allow_blank=False)
ws.add_data_validation(dv_estado)
dv_estado.add(f"I{header_row + 1}:I{last_row}")

dv_prioridad = DataValidation(type="list", formula1='"Alta,Media,Baja"', allow_blank=False)
ws.add_data_validation(dv_prioridad)
dv_prioridad.add(f"G{header_row + 1}:G{last_row}")

# =================================================================
# Hoja 2: Resumen
# =================================================================
ws2 = wb.create_sheet("Resumen")
ws2.merge_cells("A1:C1")
ws2["A1"] = "Resumen del plan de revisión"
ws2["A1"].font = title_font
ws2.row_dimensions[1].height = 22

def add_table(ws2, start_row, title, col_letter, categories, data_range):
    r = start_row
    ws2.cell(row=r, column=1, value=title).font = Font(name=FONT_NAME, bold=True, size=11)
    r += 1
    ws2.cell(row=r, column=1, value="Categoría").font = header_font
    ws2.cell(row=r, column=1).fill = header_fill
    ws2.cell(row=r, column=2, value="N° de comentarios").font = header_font
    ws2.cell(row=r, column=2).fill = header_fill
    r += 1
    first_data_row = r
    for cat in categories:
        ws2.cell(row=r, column=1, value=cat).font = base_font
        formula = f'=COUNTIF({data_range},A{r})'
        ws2.cell(row=r, column=2, value=formula).font = base_font
        ws2.cell(row=r, column=2).alignment = Alignment(horizontal="center")
        r += 1
    total_row = r
    ws2.cell(row=r, column=1, value="Total").font = Font(name=FONT_NAME, bold=True)
    ws2.cell(row=r, column=2, value=f"=SUM(B{first_data_row}:B{r-1})").font = Font(name=FONT_NAME, bold=True)
    return r + 2

data_ref_fuente = f"'Plan de revisión'!$B${header_row + 1}:$B${last_row}"
data_ref_tipo = f"'Plan de revisión'!$F${header_row + 1}:$F${last_row}"
data_ref_prioridad = f"'Plan de revisión'!$G${header_row + 1}:$G${last_row}"
data_ref_estado = f"'Plan de revisión'!$I${header_row + 1}:$I${last_row}"

next_row = add_table(ws2, 3, "Por fuente", "A", ["Editor", "Revisor 1", "Revisor 2"], data_ref_fuente)
next_row = add_table(ws2, next_row, "Por tipo de acción", "A",
                      ["Documentación", "Reanálisis", "Redacción", "Redacción / Reencuadre", "Corrección de error"],
                      data_ref_tipo)
next_row = add_table(ws2, next_row, "Por prioridad", "A", ["Alta", "Media", "Baja"], data_ref_prioridad)
next_row = add_table(ws2, next_row, "Por estado", "A",
                      ["Pendiente", "En progreso", "Completado", "No se abordará"], data_ref_estado)

ws2.column_dimensions["A"].width = 26
ws2.column_dimensions["B"].width = 18

wb.save("/home/claude/plan_revision_saberpro.xlsx")
print("saved")
